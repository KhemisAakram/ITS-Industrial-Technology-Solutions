"""Live readers — parse the ITS Profile vault files into Python data.

The folder itself is the database: no caching, each call re-reads the real files.
"""
from __future__ import annotations

import csv
import json
import re
from datetime import datetime
from pathlib import Path

import openpyxl
import yaml

BASE = Path(__file__).resolve().parent.parent
DOCS = BASE / "documents"
PLANNING = BASE / "Planning"
SERVICES = BASE / "Services"

FINANCE_CSV = DOCS / "05_Financial" / "ITS_Finance_Tracker.csv"
DAILY_DIR = PLANNING / "05_Daily_Notes"
PROJECTS_DIR = PLANNING / "04_Projects"
JOB_TRACKER_XLSX = SERVICES / "PCB Service" / "04_Tracking" / "PCB_Job_Tracker.xlsx"
PROSPECT_XLSX = DOCS / "02_Prospecting" / "ITS_Prospecting_Dashboard.xlsx"
CLIENTS_CSV = DOCS / "02_Prospecting" / "ITS_Clients.csv"
LEADS_CSV = DOCS / "02_Prospecting" / "ITS_Leads.csv"
OBSIDIAN_JSON = DOCS / "07_Time_Management" / "obsidian_data.json"
INVOICES_DIR = DOCS / "05_Financial" / "Invoices"
OPERATIONS_INVOICE = DOCS / "06_Operations" / "Invoice_3DP-003_2026-08-26.md"

FINANCE_HEADER = ["Date", "Type", "Category", "Description", "Amount (DA)"]

SKIP_DIRS = {".git", ".gitkeep"}


# --------------------------------------------------------------------------- helpers
def _num(value, default=0.0):
    if value is None:
        return default
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(",", "").replace(" DA", "").replace(" ", "")
    if not s:
        return default
    try:
        return float(s)
    except ValueError:
        return default


def _date_str(value):
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    s = str(value).strip()
    if " " in s:
        s = s.split(" ")[0]
    return s


# --------------------------------------------------------------------------- finance
def finance():
    """Parse the finance CSV -> list of dicts (index added for editing)."""
    rows = []
    if not FINANCE_CSV.exists():
        return rows
    with open(FINANCE_CSV, newline="", encoding="utf-8") as f:
        for i, r in enumerate(csv.DictReader(f)):
            rows.append({
                "index": i,
                "date": r.get("Date", "").strip(),
                "type": (r.get("Type") or "").strip(),
                "category": (r.get("Category") or "").strip(),
                "description": (r.get("Description") or "").strip(),
                "amount": _num(r.get("Amount (DA)") or r.get("Amount")),
            })
    rows.sort(key=lambda r: (r["date"], r["index"]))
    return rows


# --------------------------------------------------------------------------- daily notes
_FRONT_RE = re.compile(r"^---\s*$")
_FOCUS_RE = re.compile(r"^\s*>\s*Focus:\s*(.+?)\s*$")


def _parse_frontmatter_block(lines):
    """Return (dict, start_idx, end_idx) for first YAML frontmatter block."""
    start = end = None
    found = []
    for i, ln in enumerate(lines):
        if _FRONT_RE.match(ln):
            found.append(i)
            if len(found) == 2:
                start, end = found[0], found[1]
                break
    if start is None:
        return {}, None, None
    try:
        fm = yaml.safe_load("\n".join(lines[start + 1:end]))
    except Exception:
        fm = {}
    return (fm or {}), start, end


def _parse_md_day(path: Path):
    lines = path.read_text(encoding="utf-8").split("\n")
    fm, _, _ = _parse_frontmatter_block(lines)
    focus = ""
    for ln in lines:
        m = _FOCUS_RE.match(ln)
        if m:
            focus = m.group(1).strip()
            break
    date = fm.get("date") or path.stem
    return {
        "date": _date_str(date),
        "week": fm.get("week", ""),
        "day": fm.get("day", ""),
        "energy": fm.get("energy", ""),
        "hours_worked": _num(fm.get("hours_worked")),
        "hours_billable": _num(fm.get("hours_billable")),
        "revenue_da": _num(fm.get("revenue_da")),
        "yield_da_hr": _num(fm.get("yield_da_hr")),
        "income_da": _num(fm.get("income_da")),
        "expense_parts": _num(fm.get("expense_parts")),
        "expense_tools": _num(fm.get("expense_tools")),
        "expense_transport": _num(fm.get("expense_transport")),
        "expense_overhead": _num(fm.get("expense_overhead")),
        "status": fm.get("status", ""),
        "focus": focus,
    }


def daily_notes():
    notes = []
    if not DAILY_DIR.exists():
        return notes
    for p in sorted(DAILY_DIR.glob("????-??-??.md")):
        try:
            notes.append(_parse_md_day(p))
        except Exception:
            continue
    notes.sort(key=lambda n: n["date"])
    return notes


def obsidian_daily_fallback():
    """Fallback: read obsidian_data.json if markdown parsing yields nothing."""
    if not OBSIDIAN_JSON.exists():
        return []
    try:
        data = json.loads(OBSIDIAN_JSON.read_text(encoding="utf-8"))
        return data.get("daily", [])
    except Exception:
        return []


# --------------------------------------------------------------------------- projects
def projects():
    """Parse Planning/04_Projects/*.md frontmatter + checklist state."""
    out = []
    if not PROJECTS_DIR.exists():
        return out
    for p in sorted(PROJECTS_DIR.glob("*.md")):
        lines = p.read_text(encoding="utf-8").split("\n")
        fm, _, _ = _parse_frontmatter_block(lines)
        body = "\n".join(lines)
        tasks_done = body.count("- [x]")
        tasks_total = body.count("- [ ]") + tasks_done
        name = p.stem.replace("_", " ").replace(".md", "")
        title = name
        for ln in lines:
            if ln.startswith("# "):
                title = ln[2:].strip()
                break
        tags = fm.get("tags") or []
        if isinstance(tags, str):
            tags = [t.strip() for t in tags.strip("[]").split(",") if t.strip()]
        due = fm.get("due", "")
        progress = _num(fm.get("progress", 0), 0)
        out.append({
            "id": p.stem,
            "filename": p.name,
            "path": str(p.relative_to(BASE)),
            "title": title,
            "type": fm.get("type", ""),
            "status": fm.get("status", ""),
            "due": _date_str(due),
            "progress": int(progress) if progress == int(progress) else progress,
            "tags": tags,
            "tasks_done": tasks_done,
            "tasks_total": tasks_total,
            "tasks_pct": round(tasks_done / tasks_total * 100) if tasks_total else 0,
        })
    return out


# --------------------------------------------------------------------------- jobs
JOBS_COLS = ["Date In", "WO #", "Customer Name", "Company", "Phone", "Board Model",
             "Equipment", "Symptoms", "Diagnosis", "Quote (DA)", "Accepted?",
             "Date Out", "Paid (DA)", "Status", "Notes"]


def jobs():
    rows = []
    if not JOB_TRACKER_XLSX.exists():
        return rows
    wb = openpyxl.load_workbook(JOB_TRACKER_XLSX, data_only=True)
    if "Job Tracker" not in wb.sheetnames:
        return rows
    ws = wb["Job Tracker"]
    header = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=1, values_only=True), start=1):
        header = [str(c).strip() if c is not None else "" for c in row]
        break
    if header is None:
        return rows
    idx = {col: pos for pos, col in enumerate(JOBS_COLS) if col in header}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not any(str(c).strip() for c in r if c is not None):
            continue
        job = {col: (r[idx[col]] if col in idx else "") for col in JOBS_COLS}
        if not str(job.get("WO #") or "").strip() and not str(job.get("Customer Name") or "").strip():
            continue
        rows.append({
            "date_in": _date_str(job.get("Date In")),
            "wo": (job.get("WO #") or "").strip(),
            "customer": (job.get("Customer Name") or "").strip(),
            "company": (job.get("Company") or "").strip(),
            "phone": (job.get("Phone") or "").strip(),
            "board": (job.get("Board Model") or "").strip(),
            "equipment": (job.get("Equipment") or "").strip(),
            "symptoms": (job.get("Symptoms") or "").strip(),
            "diagnosis": (job.get("Diagnosis") or "").strip(),
            "quote": _num(job.get("Quote (DA)")),
            "accepted": (job.get("Accepted?") or "").strip(),
            "date_out": _date_str(job.get("Date Out")),
            "paid": _num(job.get("Paid (DA)")),
            "status": (job.get("Status") or "").strip(),
            "notes": (job.get("Notes") or "").strip(),
        })
    return rows


# --------------------------------------------------------------------------- clients / leads (new CSV files in the vault)
def _read_table_csv(path, cols):
    rows = []
    if not path.exists():
        return rows
    with open(path, newline="", encoding="utf-8") as f:
        for i, r in enumerate(csv.DictReader(f)):
            row = {c: (r.get(c) or "").strip() for c in cols}
            row["index"] = i
            rows.append(row)
    return rows


CLIENT_COLS = ["id", "name", "company", "phone", "email", "source", "notes"]


def clients():
    return _read_table_csv(CLIENTS_CSV, CLIENT_COLS)


LEAD_COLS = ["id", "company", "contact", "title", "phone", "zone", "status",
             "value_da", "last_contact", "notes"]


def leads():
    return _read_table_csv(LEADS_CSV, LEAD_COLS)


def leads_prospecting_xlsx():
    """Read-only view of the Skikda mapping sheet from the Prospecting dashboard."""
    rows = []
    if not PROSPECT_XLSX.exists():
        return rows
    wb = openpyxl.load_workbook(PROSPECT_XLSX, data_only=True)
    if "Ground War" not in wb.sheetnames:
        return rows
    ws = wb["Ground War"]
    for r in ws.iter_rows(min_row=5, values_only=True):
        if not r or not (r[1] or r[2]):
            continue
        rows.append({
            "company": str(r[1] or "").strip(),
            "zone": str(r[2] or "").strip(),
            "contact": str(r[4] or "").strip(),
            "status": str(r[5] or "").strip(),
            "follow_up": str(r[7] or "").strip(),
        })
    return rows


# --------------------------------------------------------------------------- invoices
def invoices():
    out = []
    if not INVOICES_DIR.exists():
        return out
    total = re.compile(r"[Tt]otal\s*\|\s*([\d, ]+)\s*\|")
    paid = re.compile(r"(?:[Pp]aido?u?t?)\s*[:#]\s*(.+)$")
    for p in sorted(INVOICES_DIR.glob("*.md")):
        text = p.read_text(encoding="utf-8")
        lines = text.split("\n")
        num = p.stem
        date = ""
        client = ""
        for ln in lines:
            m = re.match(r"\*\*(?:Invoice|N)(?: #)?\*?\*?:\s*(.+)$", ln)
            if not m and "Invoice #:" in ln:
                num = ln.split(":", 1)[1].strip()
            m = re.match(r"\*\*Date:\*\*\s*(.+)$", ln)
            if m:
                date = m.group(1).strip()
            m = re.match(r"\*\*Client:\*\*\s*(.+)$", ln)
            if m:
                client = m.group(1).strip()
        tm = total.search(text)
        amount = _num(tm.group(1)) if tm else 0
        paid_flag = False
        note = ""
        for ln in lines:
            ls = ln.strip().lower()
            if "status:** paid" in ls or ls.startswith("paid"):
                paid_flag = True
                note = ln.strip()
        out.append({
            "number": num,
            "filename": p.name,
            "path": str(p.relative_to(BASE)),
            "date": date,
            "client": client,
            "amount": amount,
            "paid": paid_flag,
            "note": note,
            "short": text[:2000],
        })
    return out


# --------------------------------------------------------------------------- documents registry
def documents():
    out = []
    for root in (DOCS, SERVICES, PLANNING):
        if not root.exists():
            continue
        for p in sorted(root.rglob("*")):
            if p.is_dir():
                continue
            if p.suffix.lower() not in (".md", ".pdf", ".docx", ".xlsx", ".html", ".csv",
                                        ".xls", ".pptx", ".png", ".jpg", ".jpeg", ".txt",
                                        ".js", ".gs", ".json", ".ps1", ".py"):
                continue
            rel = p.relative_to(BASE)
            parts = list(rel.parts)
            rel_path = str(rel)
            folder = " / ".join(parts[:-1]) if len(parts) > 1 else "root"
            out.append({
                "name": p.name,
                "ext": p.suffix.lower().lstrip("."),
                "folder": folder,
                "path": rel_path,
                "size_kb": round(p.stat().st_size / 1024, 1),
            })
    return out


def categories():
    """Distinct finance categories and types for forms."""
    cats = {}
    for t in finance():
        cats.setdefault(t["type"], set()).add(t["category"])
    return {k: sorted(v) for k, v in cats.items()}
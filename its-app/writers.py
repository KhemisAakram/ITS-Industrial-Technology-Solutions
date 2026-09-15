"""Safe writers — write edits straight back into the vault files.

Every write first copies the target to `path.bak` so a bad edit never loses data.
"""
from __future__ import annotations

import csv
import re
import shutil
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl

from readers import (BASE, CLIENT_COLS, CLIENTS_CSV, DAILY_DIR, EXPENSE_BUCKETS,
                     FINANCE_CSV, FINANCE_HEADER, INVOICES_DIR, JOB_TRACKER_XLSX,
                     LEAD_COLS, LEADS_CSV, PROJECTS_DIR)

_FRONT_RE = re.compile(r"^---\s*$")


class BackupError(Exception):
    pass


def backup(path: Path) -> None:
    """Copy `path` to `path + .bak` before any modification."""
    if not path.exists():
        return
    bak = path.with_name(path.name + ".bak")
    shutil.copy2(path, bak)


def _write_text_utf8(path: Path, content: str) -> None:
    backup(path)
    path.write_text(content, encoding="utf-8")


def _rel(p: Path) -> str:
    return str(p.relative_to(BASE))


# --------------------------------------------------------------------------- finance
def write_finance(rows) -> None:
    """Rewrite ITS_Finance_Tracker.csv from a list of dicts (index ignored)."""
    ordered = sorted(rows, key=lambda r: (r["date"], r["description"]))
    lines = [",".join(FINANCE_HEADER)]
    for r in ordered:
        def esc(s):
            s = str(s)
            if "," in s or '"' in s or "\n" in s:
                return '"' + s.replace('"', '""') + '"'
            return s
        lines.append(",".join([
            esc(r["date"]), esc(r["type"]), esc(r["category"]),
            esc(r["description"]), esc(r["amount"]),
        ]))
    _write_text_utf8(FINANCE_CSV, "\n".join(lines) + "\n")


def add_finance(entry) -> int:
    rows = _read_finance_rows()
    row = {
        "date": entry.get("date", ""),
        "type": entry.get("type", "income"),
        "category": entry.get("category", ""),
        "description": entry.get("description", ""),
        "amount": entry.get("amount", 0),
    }
    rows.append(row)
    write_finance(rows)
    return len(rows) - 1


def _read_finance_rows():
    import readers
    return [dict(r) for r in readers.finance()]


def update_finance(index, entry) -> None:
    import readers
    rows = [dict(r) for r in readers.finance()]
    if not (0 <= index < len(rows)):
        return
    for k, v in entry.items():
        rows[index][k] = v
    write_finance(rows)


def delete_finance(index) -> None:
    import readers
    rows = [dict(r) for r in readers.finance()]
    if 0 <= index < len(rows):
        rows.pop(index)
        write_finance(rows)


# --------------------------------------------------------------------------- frontmatter (daily notes / projects)
def _replace_frontmatter(path: Path, fields: dict) -> None:
    """Surgically replace or add YAML frontmatter keys, preserving the rest of the file."""
    lines = path.read_text(encoding="utf-8").split("\n")
    markers = [i for i, ln in enumerate(lines) if _FRONT_RE.match(ln)]
    if len(markers) < 2:
        raise ValueError(f"No frontmatter in {path}")
    start, end = markers[0], markers[1]

    key_re = {k: re.compile(r"^" + re.escape(k) + r"(\s*:\s*)(.*)$") for k in fields}
    seen = set()
    out = lines[: start + 1]
    inserted = []
    for ln in lines[start + 1: end]:
        matched = False
        for k, rx in key_re.items():
            m = rx.match(ln)
            if m and k not in seen:
                out.append(f"{k}: {str(fields[k]).strip()}")
                seen.add(k)
                matched = True
                break
        if not matched:
            out.append(ln)
    missing = [k for k in fields if k not in seen]
    if missing:
        for k in missing:
            inserted.append(f"{k}: {fields[k]}")
        if inserted:
            out = out[: start + 1] + inserted + out[start + 1:]
    out.append(lines[end])
    for ln in lines[end + 1:]:
        out.append(ln)
    _write_text_utf8(path, "\n".join(out))


def update_daily_note(date: str, fields: dict) -> None:
    path = DAILY_DIR / f"{date}.md"
    if not path.exists():
        raise FileNotFoundError(f"Daily note not found: {path}")
    _replace_frontmatter(path, fields)


DAILY_TAG = "[daily "


def _is_daily_tagged(row) -> bool:
    return str(row.get("description", "")).startswith(DAILY_TAG)


def _bucket_of(row):
    """Map a finance row to a daily-note bucket key."""
    if row["type"] == "income":
        return ("income", "*")
    cat = (row.get("category") or "").strip().lower()
    if cat in EXPENSE_BUCKETS:
        return ("expense", cat)
    return ("expense", "other")


def sync_daily_to_finance(date: str):
    """Push daily note totals into the finance CSV (reconcile, no duplicates).

    Only adds a tagged '[daily <date>] <bucket>' row when the note amount is
    greater than what is already logged for the same bucket on that date.
    Tagged rows for this date are refreshed first. Returns created entries.
    """
    import readers
    note = readers.daily_note(date)
    if not note:
        return []
    target = {
        ("income", "*"): note["income_da"],
        ("expense", "parts"): note["expense_parts"],
        ("expense", "tools"): note["expense_tools"],
        ("expense", "transport"): note["expense_transport"],
        ("expense", "overhead"): note["expense_overhead"],
        ("expense", "other"): note["expense_other"],
    }

    rows = [dict(r) for r in readers.finance()]
    day_rows = [r for r in rows if r["date"] == date]
    kept = [r for r in rows if not (_is_daily_tagged(r) and r["date"] == date)]
    removed = len(kept) != len(rows)
    rows = kept

    manual = {}
    for r in day_rows:
        if _is_daily_tagged(r):
            continue
        key = _bucket_of(r)
        manual[key] = manual.get(key, 0) + r.get("amount", 0)

    created = []
    for (typ, cat), amt in target.items():
        if amt <= 0:
            continue
        existing = manual.get((typ, cat), 0)
        delta = amt - existing
        if delta > 0:
            word = "income" if typ == "income" else cat
            row = {"date": date, "type": typ,
                   "category": "General" if typ == "income"
                   else (cat if cat in EXPENSE_BUCKETS else "Other"),
                   "description": f"{DAILY_TAG}{date}] {word}", "amount": delta}
            rows.append(row)
            created.append(row)

    if created or removed:
        write_finance(rows)
    return created


def _clean_num(v):
    v = float(v or 0)
    return int(v) if v == int(v) else v


def sync_finance_to_daily(date: str) -> bool:
    """Pull finance totals for `date` back into the daily note's money fields.

    Makes the daily note equal the finance ledger for that day:
    income_da = all income; each expense bucket = that category's total;
    expense_other = any expense not in the standard buckets. Returns True if
    the daily note was changed.
    """
    import readers
    note = readers.daily_note(date)
    if not note:
        return False

    rows = [r for r in readers.finance() if r["date"] == date]
    income = sum(r["amount"] for r in rows if r["type"] == "income")
    buckets = {"parts": 0.0, "tools": 0.0, "transport": 0.0, "overhead": 0.0, "other": 0.0}
    for r in rows:
        if r["type"] != "expense":
            continue
        key = _bucket_of(r)[1]
        buckets[key] += r["amount"]

    fields = {
        "income_da": _clean_num(income),
        "expense_parts": _clean_num(buckets["parts"]),
        "expense_tools": _clean_num(buckets["tools"]),
        "expense_transport": _clean_num(buckets["transport"]),
        "expense_overhead": _clean_num(buckets["overhead"]),
        "expense_other": _clean_num(buckets["other"]),
    }
    changed = any(abs((note.get(k) or 0) - float(v)) > 0.001 for k, v in fields.items())
    if changed:
        update_daily_note(date, fields)
    return changed


# --------------------------------------------------------------------------- projects
def update_project(project_id: str, fields: dict) -> None:
    path = PROJECTS_DIR / f"{project_id}.md"
    if not path.exists():
        raise FileNotFoundError(f"Project not found: {path}")
    _replace_frontmatter(path, fields)


# --------------------------------------------------------------------------- jobs (XLSX)
JOBS_COLS = ["Date In", "WO #", "Customer Name", "Company", "Phone", "Board Model",
             "Equipment", "Symptoms", "Diagnosis", "Quote (DA)", "Accepted?",
             "Date Out", "Paid (DA)", "Status", "Notes"]


def _load_job_workbook():
    wb = openpyxl.load_workbook(JOB_TRACKER_XLSX)  # keep formulas/styles
    return wb


def write_jobs(job_rows) -> None:
    backup(JOB_TRACKER_XLSX)
    wb = _load_job_workbook()
    if "Job Tracker" not in wb.sheetnames:
        raise ValueError("Job Tracker sheet missing")
    ws = wb["Job Tracker"]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for row in job_rows:
        ws.append([
            row.get("date_in", ""), row.get("wo", ""), row.get("customer", ""),
            row.get("company", ""), row.get("phone", ""), row.get("board", ""),
            row.get("equipment", ""), row.get("symptoms", ""), row.get("diagnosis", ""),
            row.get("quote", 0), row.get("accepted", ""), row.get("date_out", ""),
            row.get("paid", 0), row.get("status", "Received"), row.get("notes", ""),
        ])
    wb.save(JOB_TRACKER_XLSX)


# --------------------------------------------------------------------------- clients / leads (CSV)
def write_clients(rows) -> None:
    _write_csv(CLIENTS_CSV, CLIENT_COLS, rows)


def write_leads(rows) -> None:
    _write_csv(LEADS_CSV, LEAD_COLS, rows)


def _write_csv(path: Path, cols: list, rows) -> None:
    lines = [",".join(cols)]
    for r in rows:
        cells = []
        for c in cols:
            v = str(r.get(c, "") or "")
            if "," in v or '"' in v or "\n" in v:
                v = '"' + v.replace('"', '""') + '"'
            cells.append(v)
        lines.append(",".join(cells))
    _write_text_utf8(path, "\n".join(lines) + "\n")


# --------------------------------------------------------------------------- invoices
INVOICE_TAG = "[invoice "


def _is_invoice_tagged(row) -> bool:
    return str(row.get("description", "")).startswith(INVOICE_TAG)


def sync_invoice_to_finance(number: str, keep=True):
    """Reconcile finance rows for one invoice.

    Removes any '[invoice <number>]' income row, then — if the invoice is
    Paid and keep=True — re-adds a single tagged row so income is never
    double counted. Returns list of created rows.
    """
    import readers
    inv = readers.invoice_by_number(number)
    rows = [dict(r) for r in readers.finance()]
    tag = f"{INVOICE_TAG}{number}]"
    kept = [r for r in rows if not str(r.get("description", "")).startswith(tag)]
    removed = len(kept) != len(rows)
    rows = kept
    created = []
    if inv and inv["paid"] and inv["amount"] > 0 and keep:
        date = inv["date"] or datetime.now().strftime("%Y-%m-%d")
        desc = tag if not inv["client"] else f"{tag} {inv['client']}"
        row = {"date": date, "type": "income", "category": "General",
               "description": desc, "amount": inv["amount"]}
        rows.append(row)
        created.append(row)
    if created or removed:
        write_finance(rows)
    return created


def _is_num(s) -> bool:
    try:
        float(str(s).replace(",", ""))
        return True
    except (ValueError, TypeError):
        return False


def _num(s, default=0):
    try:
        return float(str(s).replace(" ", "").replace(",", "")) if str(s).strip() else float(default)
    except (ValueError, TypeError):
        return float(default)


def parse_line_items(raw: str):
    """Parse 'description, qty, rate' lines into line-item dicts."""
    items = []
    for ln in (raw or "").splitlines():
        ln = ln.strip().rstrip("|")
        if not ln:
            continue
        parts = [p.strip() for p in re.split(r"[,|;]", ln) if p.strip()]
        if not parts:
            continue
        desc = parts[0]
        qty = int(_num(parts[1], 1)) if len(parts) > 1 and _is_num(parts[1]) else 1
        unit = "pcs"
        rate = _num(parts[2]) if len(parts) > 2 and _is_num(parts[2]) else 0
        if len(parts) == 2 and not _is_num(parts[1]):
            desc = parts[0]
            qty, unit, rate = 1, parts[1], 0
        if len(parts) > 3 and not _is_num(parts[1]):
            qty, unit = 1, parts[1]
            rate = _num(parts[2]) if len(parts) > 2 else 0
        items.append({
            "description": desc, "qty": qty, "unit": unit,
            "rate": rate, "amount": round(qty * rate, 2),
        })
    return items


def _invoice_body(number: str, client: str, date: str, amount, line_items=None,
                  status_note: str = "") -> str:
    content = (
        f"# ITS — Invoice\n\n"
        f"**Invoice #:** {number}\n"
        f"**Date:** {date}\n"
        f"**Due:** On receipt\n\n"
        f"---\n\n## Bill To\n\n"
        f"**Client:** {client or ''}\n"
    )
    if status_note:
        content += f"**Status:** {status_note}\n"
    content += "\n---\n\n## Line Items\n\n"
    content += (f"| # | Description | Qty | Unit | Rate (DA) | Amount (DA) |\n"
                f"|---|-------------|-----|------|-----------|-------------|\n")
    if line_items:
        for i, it in enumerate(line_items, start=1):
            content += (f"| {i} | {it.get('description', '')} | {it.get('qty', 1)} | "
                        f"{it.get('unit', 'pcs')} | {it.get('rate', 0)} | {it.get('amount', 0)} |\n")
        total = sum(it.get("amount", 0) for it in line_items)
    else:
        total = amount
    content += (
        f"\n---\n\n## Summary\n\n"
        f"| | Amount (DA) |\n|---|-------------|\n"
        f"| **Total** | **{int(round(total)):,}** |\n\n"
        f"---\n\n## Notes\n\n"
        f"- All prices in DA (Algerian Dinar)\n"
        f"- TVA 19% included\n\n"
        f"---\n\n**ITS — Industrial Technology Solutions**\n"
        f"El Harrouch, Skikda\n"
    )
    return content


def create_invoice(number: str, client: str, date: str, amount, line_items=None) -> Path:
    INVOICES_DIR.mkdir(parents=True, exist_ok=True)
    path = INVOICES_DIR / f"{number}.md"
    if path.exists():
        raise FileExistsError(f"Invoice {number} already exists")
    today = date or datetime.now().strftime("%Y-%m-%d")
    path.write_text(_invoice_body(number, client, today, amount, line_items),
                    encoding="utf-8")
    return path


def update_invoice(number: str, *, client=None, date=None, amount=None, line_items=None) -> None:
    """Rewrite an invoice preserving its Paid status note."""
    import readers
    inv = readers.invoice_by_number(number)
    if not inv:
        raise FileNotFoundError(f"Invoice {number} not found")
    status_note = ""
    if inv["paid"]:
        status_note = "Paid"
        if inv["note"]:
            m = re.search(r"Paid\s*\((.+?)\)", inv["note"])
            if m:
                status_note = f"Paid ({m.group(1)})"
    body = _invoice_body(
        number,
        client if client is not None else inv["client"],
        date if date is not None else inv["date"],
        amount if amount is not None else inv["amount"],
        line_items if line_items is not None else inv["line_items"],
        status_note,
    )
    path = INVOICES_DIR / f"{number}.md"
    backup(path)
    path.write_text(body, encoding="utf-8")


def toggle_invoice_paid(number: str) -> None:
    import readers
    inv = readers.invoice_by_number(number)
    if not inv:
        raise FileNotFoundError(f"Invoice {number} not found")
    path = INVOICES_DIR / f"{number}.md"
    text = path.read_text(encoding="utf-8")
    backup(path)
    if not inv["paid"]:
        date = datetime.now().strftime("%Y-%m-%d")
        if re.search(r"(?mi)^\*\*Status:\*\*", text):
            text = re.sub(r"(?mi)^\*\*Status:\*\*.*$", f"**Status:** Paid ({date})", text, count=1)
        else:
            text = re.sub(r"(?m)^(\*\*Client:\*\*.*)$", r"\1\n**Status:** Paid (" + date + ")", text, count=1)
    else:
        text = re.sub(r"(?mi)^\*\*Status:\*\*.*$\n?", "", text, count=1)
    path.write_text(text, encoding="utf-8")


def delete_invoice(number: str) -> None:
    path = INVOICES_DIR / f"{number}.md"
    if path.exists():
        backup(path)
        path.unlink()


# --------------------------------------------------------------------------- backup
BACKUP_DIR = Path(__file__).resolve().parent / "backups"
_BACKUP_NOISE = {".git", "__pycache__", "backups", ".venv", "venv", "node_modules"}


def create_backup(max_keep: int = 10) -> Path:
    """Zip the whole ITS Profile (minus noise), keep a dated copy on disk."""
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    zpath = BACKUP_DIR / f"ITS_Backup_{stamp}.zip"
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as zf:
        for p in BASE.rglob("*"):
            if p.is_dir() or p.suffix == ".bak":
                continue
            parts = p.relative_to(BASE).parts
            if any(part in _BACKUP_NOISE for part in parts):
                continue
            zf.write(p, str(p.relative_to(BASE)).replace("\\", "/"))
    prev = sorted(BACKUP_DIR.glob("ITS_Backup_*.zip"), reverse=True)
    for old in prev[max_keep:]:
        try:
            old.unlink()
        except OSError:
            pass
    return zpath


# --------------------------------------------------------------------------- next number helper
def next_invoice_number() -> str:
    from readers import invoices
    year = datetime.now().year
    prefix = f"ITS-{year}-"
    existing = [inv["number"] for inv in invoices() if inv["number"].startswith(prefix)]
    m = 0
    for num in existing:
        try:
            m = max(m, int(num.rsplit("-", 1)[1]))
        except (ValueError, IndexError):
            continue
    return f"{prefix}{m + 1:03d}"